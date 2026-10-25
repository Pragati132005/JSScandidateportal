"""Inspect raw extracted text from sample PDFs and DOCXs."""

import os
import docx
import pdfplumber
import fitz  # PyMuPDF


def inspect_all():
    inspect_dir = os.path.join("data", "_inspect")
    os.makedirs(inspect_dir, exist_ok=True)

    resumes_pdf_dir = os.path.join("data", "resumes", "PDF")
    resumes_docx_dir = os.path.join("data", "resumes", "Word")

    for fname in os.listdir(resumes_pdf_dir):
        if not fname.endswith(".pdf"):
            continue
        pdf_path = os.path.join(resumes_pdf_dir, fname)
        base = os.path.splitext(fname)[0]

        # pdfplumber
        plumber_text = []
        with pdfplumber.open(pdf_path) as pdf:
            for i, p in enumerate(pdf.pages):
                plumber_text.append(f"--- PAGE {i+1} ---")
                plumber_text.append(p.extract_text() or "")
        plumber_out = "\n".join(plumber_text)
        with open(os.path.join(inspect_dir, f"{base}_plumber.txt"), "w", encoding="utf-8") as f:
            f.write(plumber_out)

        # PyMuPDF
        fitz_text = []
        doc = fitz.open(pdf_path)
        for i, page in enumerate(doc):
            fitz_text.append(f"--- PAGE {i+1} ---")
            fitz_text.append(page.get_text() or "")
        fitz_out = "\n".join(fitz_text)
        with open(os.path.join(inspect_dir, f"{base}_fitz.txt"), "w", encoding="utf-8") as f:
            f.write(fitz_out)

    for fname in os.listdir(resumes_docx_dir):
        if not fname.endswith(".docx"):
            continue
        docx_path = os.path.join(resumes_docx_dir, fname)
        base = os.path.splitext(fname)[0]
        doc = docx.Document(docx_path)
        docx_lines = []
        for p in doc.paragraphs:
            if p.text:
                docx_lines.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                docx_lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        docx_out = "\n".join(docx_lines)
        with open(os.path.join(inspect_dir, f"{base}_docx.txt"), "w", encoding="utf-8") as f:
            f.write(docx_out)

    print("Extracted all samples to data/_inspect/")


if __name__ == "__main__":
    inspect_all()
