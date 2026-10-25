"""Tests for shared/ingest.py covering all 8 sample resumes and edge cases."""

import os
import pytest
from shared.ingest import extract_text, detect_file_format, IngestResult
from resume_parser.segment import segment_resume
from shared.clean import clean_text
from shared.llm_client import LLMResponse


CANDIDATES = [
    ("Ananya Rao", "Ananya Rao"),
    ("Karthik Subramaniam", "Karthik Subramaniam"),
    ("Priya Menon", "Priya Menon"),
    ("Rohan Verma", "Rohan Verma"),
]


@pytest.mark.parametrize("candidate_name,expected_substr", CANDIDATES)
def test_extract_pdf_samples(candidate_name, expected_substr):
    """Test PDF extraction for each candidate."""
    pdf_path = os.path.join("data", "resumes", "PDF", f"{candidate_name}.pdf")
    assert os.path.exists(pdf_path), f"File {pdf_path} must exist"

    result = extract_text(pdf_path)
    assert isinstance(result, IngestResult)
    assert result.input_format == "pdf"
    assert result.pages >= 1
    assert expected_substr in result.text
    assert len(result.text.strip()) > 100


@pytest.mark.parametrize("candidate_name,expected_substr", CANDIDATES)
def test_extract_docx_samples(candidate_name, expected_substr):
    """Test DOCX extraction for each candidate."""
    docx_path = os.path.join("data", "resumes", "Word", f"{candidate_name}.docx")
    assert os.path.exists(docx_path), f"File {docx_path} must exist"

    result = extract_text(docx_path)
    assert isinstance(result, IngestResult)
    assert result.input_format == "docx"
    assert expected_substr in result.text
    assert len(result.text.strip()) > 100


def test_detect_unsupported_format(tmp_path):
    """Verify unsupported file format raises ValueError."""
    fake_file = tmp_path / "test.txt"
    fake_file.write_text("Hello world", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file format"):
        detect_file_format(str(fake_file))


def test_nonexistent_file_raises():
    """Verify non-existent path raises FileNotFoundError."""
    with pytest.raises(FileNotFoundError):
        extract_text("non_existent_file.pdf")


def test_extract_pdf_two_column_reading_order(tmp_path):
    """Read a full-width header, then the left column, then the right column."""
    from reportlab.pdfgen.canvas import Canvas

    pdf_path = tmp_path / "two-column.pdf"
    canvas = Canvas(str(pdf_path), pagesize=(612, 792))
    canvas.drawString(50, 750, "Example Candidate")
    canvas.drawString(390, 750, "candidate@example.com")
    canvas.drawString(50, 690, "WORK EXPERIENCE")
    canvas.drawString(50, 670, "Built internal tools")
    canvas.drawString(50, 650, "Maintained deployment scripts")
    canvas.drawString(360, 690, "EDUCATION")
    canvas.drawString(360, 670, "B.S. Computer Science")
    canvas.drawString(360, 650, "Example University")
    canvas.drawString(360, 610, "SKILLS")
    canvas.drawString(360, 590, "Python")
    canvas.save()

    text = extract_text(str(pdf_path)).text
    positions = [
        text.index("Example Candidate"),
        text.index("WORK EXPERIENCE"),
        text.index("Built internal tools"),
        text.index("EDUCATION"),
        text.index("B.S. Computer Science"),
        text.index("SKILLS"),
        text.index("Python"),
    ]
    assert positions == sorted(positions)


def test_two_column_pdf_job_title_does_not_contaminate_education(tmp_path):
    """Keep a nearby right-column job title out of the education section."""
    from reportlab.pdfgen.canvas import Canvas

    pdf_path = tmp_path / "education-near-job-title.pdf"
    canvas = Canvas(str(pdf_path), pagesize=(612, 792))
    canvas.drawString(50, 750, "Synthetic Candidate")
    canvas.drawString(390, 750, "candidate@example.com")
    canvas.drawString(50, 690, "EDUCATION")
    canvas.drawString(50, 670, "B.S. Computer Science")
    canvas.drawString(390, 690, "EXPERIENCE")
    canvas.drawString(390, 670, "Software Engineer Intern")
    canvas.save()

    from resume_parser.pipeline import parse_resume

    result = parse_resume(str(pdf_path), use_llm=False)

    assert result.profile_prefill.education == "B.S., Computer Science"
    assert result.profile_prefill.education != "Software Engineer Intern"
    assert result.profile_prefill.name == "Synthetic Candidate"
    assert result.profile_prefill.email == "candidate@example.com"


def test_split_header_name_is_joined_in_rule_based_parse(tmp_path):
    """Join consecutive large-font name lines before contact details."""
    from reportlab.pdfgen.canvas import Canvas
    from resume_parser.pipeline import parse_resume

    pdf_path = tmp_path / "split-name.pdf"
    canvas = Canvas(str(pdf_path), pagesize=(612, 792))
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(50, 750, "Morgan")
    canvas.drawString(50, 725, "Reed")
    canvas.setFont("Helvetica", 10)
    canvas.drawString(50, 700, "morgan.reed@example.com")
    canvas.drawString(50, 685, "(000) 000-0000")
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(50, 650, "EDUCATION")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(50, 630, "Bachelor of Science in Computer Science")
    canvas.drawString(50, 612, "Example University")
    canvas.drawString(50, 594, "2025")
    canvas.save()

    result = parse_resume(str(pdf_path), use_llm=False)

    assert result.profile_prefill.name == "Morgan Reed"
    assert result.profile_prefill.email == "morgan.reed@example.com"
    assert result.details.education_details[0].degree == "Bachelor of Science"
    assert result.details.education_details[0].field == "Computer Science"


def test_mocked_llm_short_name_cannot_replace_font_name(tmp_path):
    """Fuse a short mocked LLM name with a stacked PDF header name."""
    from reportlab.pdfgen.canvas import Canvas
    from resume_parser.pipeline import parse_resume

    class ShortNameClient:
        def generate(self, system_prompt, user_prompt, prompt_version, cache_key_override=None):
            raw = {
                "name": "Morgan",
                "email": "morgan.reed@example.com",
                "skills": [],
                "education": [],
            }
            return LLMResponse(
                content="{}",
                parsed_json=raw,
                model_used="mock-model",
                provider_used="mock",
                from_cache=False,
            )

    pdf_path = tmp_path / "mocked-short-name.pdf"
    canvas = Canvas(str(pdf_path), pagesize=(612, 792))
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(50, 750, "Morgan")
    canvas.drawString(50, 725, "Reed")
    canvas.setFont("Helvetica", 10)
    canvas.drawString(50, 700, "morgan.reed@example.com")
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(50, 650, "EDUCATION")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(50, 630, "Bachelor of Science in Computer Science")
    canvas.save()

    result = parse_resume(str(pdf_path), use_llm=True, client=ShortNameClient())

    assert result.profile_prefill.name == "Morgan Reed"



def test_unseen_two_column_resume_sections_are_detected():
    """Verify the unseen two-column resume keeps each section contiguous."""
    path = os.path.join(
        "data", "unseen", "computer-science-internship-resume-example_3.pdf"
    )
    if not os.path.exists(path):
        pytest.skip("Organizer unseen-resume pack is not present")
    text = extract_text(path).text
    segmented = segment_resume(clean_text(text).text)

    assert {"experience", "education", "skills"}.issubset(segmented.sections)
    assert "Technician" in segmented.sections["experience"]
    assert "B.S." in segmented.sections["education"]
    assert "Java" in segmented.sections["skills"]
    assert "EDUCATION" not in segmented.sections["experience"]
    assert "SKILLS" not in segmented.sections["education"]


def test_extract_docx_tables_in_document_order(tmp_path):
    """Keep table rows and cells in place relative to surrounding paragraphs."""
    from docx import Document

    docx_path = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    document.add_paragraph("After table")
    document.save(str(docx_path))

    text = extract_text(str(docx_path)).text
    assert text.index("Before table") < text.index("A1 | B1")
    assert text.index("A1 | B1") < text.index("A2 | B2")
    assert text.index("A2 | B2") < text.index("After table")
